"""python -m unittest test_main   (with requirements.txt installed)"""

import io
import unittest

import pymupdf
from docx import Document

import main


# The built-in font is Latin-1 only; Turkish was checked on real files.
def text_pdf(lines=("Mietvertrag für Büro", "Madde 1: Taraflar"), bullet="-") -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for line in lines:
        page.insert_text((72, y), line, fontsize=12)
        y += 20
    page.insert_text((72, y), f"{bullet} first point of the list", fontsize=12)
    return doc.tobytes()


def scan_pdf(ocr_layer: bool) -> bytes:
    """A page that is a picture, optionally with invisible OCR text on it."""
    picture = pymupdf.open(stream=text_pdf(lines=["Scanned page"] * 20), filetype="pdf")
    pix = picture[0].get_pixmap(dpi=72)
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, pixmap=pix)
    if ocr_layer:
        for i in range(20):
            page.insert_text((72, 72 + 20 * i), "Scanned page text", render_mode=3)
    return doc.tobytes()


class CheckTest(unittest.TestCase):
    def refused(self, pdf):
        with self.assertRaises(main.Refused) as ctx:
            main.check(pdf)
        return ctx.exception.status, ctx.exception.code

    def test_refuses_what_is_not_a_usable_pdf(self):
        self.assertEqual(self.refused(b""), (400, "empty"))
        self.assertEqual(self.refused(b"PK\x03\x04 a zip"), (400, "not_pdf"))
        self.assertEqual(self.refused(b"%PDF-1.7 broken"), (400, "not_pdf"))

    def test_scans_are_refused_with_or_without_an_ocr_layer(self):
        self.assertEqual(self.refused(scan_pdf(ocr_layer=False)), (422, "scanned"))
        self.assertEqual(self.refused(scan_pdf(ocr_layer=True)), (422, "scanned"))

    def test_encrypted_and_oversized(self):
        doc = pymupdf.open(stream=text_pdf(), filetype="pdf")
        locked = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")
        self.assertEqual(self.refused(locked), (400, "encrypted"))
        many = pymupdf.open()
        for _ in range(main.MAX_PAGES + 1):
            many.new_page().insert_text((72, 72), "page with some visible text " * 3)
        self.assertEqual(self.refused(many.tobytes()), (413, "too_many_pages"))


class ConvertTest(unittest.TestCase):
    def test_text_pdf_becomes_an_editable_docx(self):
        out = main.convert(text_pdf())
        text = "\n".join(p.text for p in Document(io.BytesIO(out)).paragraphs)
        self.assertIn("Mietvertrag für Büro", text)
        self.assertIn("Madde 1: Taraflar", text)

    def test_private_use_bullets_become_real_ones(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("\uf0b7").font.name = "Symbol"
        p.add_run(" item")
        raw = io.BytesIO()
        doc.save(raw)
        fixed = Document(io.BytesIO(main.fix_bullets(raw.getvalue())))
        run = fixed.paragraphs[0].runs[0]
        self.assertEqual(run.text, "\u2022")
        self.assertEqual(run.font.name, "Arial")

    def test_http_errors_carry_the_code(self):
        client = main.app.test_client()
        r = client.post("/convert", data=scan_pdf(ocr_layer=True))
        self.assertEqual(r.status_code, 422)
        self.assertEqual(r.get_json(), {"error": "scanned"})
        r = client.post("/convert", data=text_pdf())
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data.startswith(b"PK"))


if __name__ == "__main__":
    unittest.main()
